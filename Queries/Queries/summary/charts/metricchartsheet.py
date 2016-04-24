import logging
from   xlinterface.xlworkbook     import XlWorkBook
from   xlinterface.xlworksheet    import XlWorkSheet
from   summary.matrix.matrixdata  import MatrixData 
from   summary.matrix.matrixtable import MatrixTable
from   summary.charts.chartdata   import ChartData
from   openpyxl.chart             import LineChart
from   openpyxl.chart             import BarChart
from   openpyxl.chart.reference   import Reference

#----------------------------------------------------------------------
class MetricChartSheet:
  #--------------------------------------------------------------------
  def __init__(self,ws,region,type,period):
    self.ws     = ws
    self.region = region
    self.type   = type
    self.period = period

    startRow = 200
    startCol = 200

    self.tables = {}

    data  = MatrixData(region,'ACTIVITY',period)
    table = MatrixTable(ws,startRow,startCol,data)
    self.tables['ACTIVITY'] = (data,table)

    startRow += data.table.dataRows + 2

    data  = MatrixData(region,'LTS',period)
    table = MatrixTable(ws,startRow,startCol,data)
    self.tables['LTS'] = (data,table)

    startRow += data.table.dataRows + 2

    data  = MatrixData(region,'UTL-CF',period)
    table = MatrixTable(ws,startRow,startCol,data)
    self.tables['ULT-CF'] = (data,table)

    startRow += data.table.dataRows + 2

    data  = MatrixData(region,'UTL-PS',period)
    table = MatrixTable(ws,startRow,startCol,data)
    self.tables['ULT-PF'] = (data,table)

    startRow += data.table.dataRows + 2

    data  = MatrixData(region,'UTL-DT',period)
    table = MatrixTable(ws,startRow,startCol,data)
    self.tables['ULT-DT'] = (data,table)

    startRow += data.table.dataRows + 2

    data  = MatrixData(region,'UTL-LS',period)
    table = MatrixTable(ws,startRow,startCol,data)
    self.tables['ULT-LS'] = (data,table)

    startRow += data.table.dataRows + 2

    data  = MatrixData(region,'OVERTIME',period)
    table = MatrixTable(ws,startRow,startCol,data)
    self.tables['OVERTIME'] = (data,table)

    startRow += data.table.dataRows + 2

    data  = MatrixData(region,'GKA',period)
    table = MatrixTable(ws,startRow,startCol,data)
    self.tables['GKA'] = (data,table)

    startRow += data.table.dataRows + 2

    actList = [10,11,12,13,14,15,16,17,18,19,20,21,22,23]
    locList = ['UK','Sweden','Finland','France','Germany','Other (EMEA)']
    for act in actList:
      data  = MatrixData(region,'ACT-BY-LOC',period,act=act,loc=locList)
      table = MatrixTable(ws,startRow,startCol,data)
      self.tables['ACT-' + str(act)] = (data,table)

      startRow += data.table.dataRows + 2
    logging.debug('')

    self.drawUtlChart()
    self.drawActChart()

  #--------------------------------------------------------------------
  def drawUtlChart(self):
    chart = LineChart()
    chart.title = 'Util'

    row = 224
    col = self.ws.GetColumnIndex('GS')
    data = Reference(self.ws.ws,min_row=row,min_col=col,max_row=row,max_col=col+13)
    chart.add_data(data,from_rows=True)

    row = 229
    col = self.ws.GetColumnIndex('GS')
    data = Reference(self.ws.ws,min_row=row,min_col=col,max_row=row,max_col=col+13)
    chart.add_data(data,from_rows=True)

    row = 234
    col = self.ws.GetColumnIndex('GS')
    data = Reference(self.ws.ws,min_row=row,min_col=col,max_row=row,max_col=col+13)
    chart.add_data(data,from_rows=True)

    row = 239
    col = self.ws.GetColumnIndex('GS')
    data = Reference(self.ws.ws,min_row=row,min_col=col,max_row=row,max_col=col+13)
    chart.add_data(data,from_rows=True)

    self.ws.ws.add_chart(chart,'B2')

  #--------------------------------------------------------------------
  def drawActChart(self):

    chart = BarChart()
    chart.title = 'Activites'

    row1 = 200
    row2 = 214
    col1 = self.ws.GetColumnIndex('GR')
    col2 = self.ws.GetColumnIndex('HF')

    data = Reference(self.ws.ws,min_row=row1,min_col=col1,max_row=row2,max_col=col2)
    chart.add_data(data)

    chart.type     = 'col'
    chart.style    =   12
    chart.grouping = 'stacked'
    chart.overlap  = 100

    self.ws.ws.add_chart(chart,'B17')

